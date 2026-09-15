from __future__ import annotations

from copy import deepcopy
from io import BytesIO
from pathlib import Path
import sys
import os
import tempfile
import unittest
from zipfile import ZipFile

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))


class OnlineE2EPlanTests(unittest.TestCase):
    def setUp(self):
        from online_e2e.contracts import load_target_config

        self.target = load_target_config("data-law-case-expert", root=ROOT)

    def _online_config(self):
        skill_info = []
        for name in self.target.expected_skill_order:
            skill_info.append(
                {
                    "skillNid": self.target.online_skill_nids[name],
                    "name": name,
                    "enabled": True,
                    "permission": "TEACHER",
                    "bindingSource": self.target.online_skill_binding_sources[name],
                    "serverMetadata": {"preserved": name},
                }
            )
        return {
            "basicInfo": {"nid": self.target.expert_nid, "isPublish": 1},
            "expertMd": {
                "templateNid": "eu5FW3sdWS",
                "customContent": "old",
                "rawContent": "keep-raw",
            },
            "agentMd": None,
            "skillInfoList": skill_info,
            "datasets": None,
            "serverOnly": {"must": "survive"},
        }

    def test_desired_config_clones_server_config_and_only_replaces_expert_content_and_skill_order(self):
        from online_e2e.desired_config import build_desired_config

        current = self._online_config()
        original = deepcopy(current)
        desired = build_desired_config(current, self.target, "new local Agent.md")

        self.assertEqual(current, original)
        self.assertEqual(desired["expertMd"]["customContent"], "new local Agent.md")
        self.assertEqual(desired["expertMd"]["rawContent"], "keep-raw")
        self.assertEqual(desired["serverOnly"], {"must": "survive"})
        self.assertEqual(
            [item["name"] for item in desired["skillInfoList"]],
            list(self.target.expected_skill_order),
        )
        self.assertEqual(
            desired["skillInfoList"][0]["serverMetadata"],
            {"preserved": "search-router"},
        )

    def test_desired_config_rejects_binding_source_different_from_target(self):
        from online_e2e.desired_config import DesiredConfigError, build_desired_config

        current = self._online_config()
        current["skillInfoList"][0]["bindingSource"] = "MARKETPLACE"
        with self.assertRaisesRegex(DesiredConfigError, "binding_source_mismatch"):
            build_desired_config(current, self.target, "new")

    def test_desired_and_synthetic_share_stable_recursive_json_clone(self):
        from online_e2e.json_clone import clone_json

        source = {"nested": {"items": (1, "two", None)}}
        cloned = clone_json(source)
        self.assertEqual(cloned, {"nested": {"items": [1, "two", None]}})
        self.assertIsNot(cloned, source)
        desired_source = (ROOT / "online_e2e" / "desired_config.py").read_text(encoding="utf-8")
        synthetic_source = (ROOT / "online_e2e" / "synthetic_backend.py").read_text(encoding="utf-8")
        self.assertIn("from .json_clone import clone_json", desired_source)
        self.assertIn("from .json_clone import clone_json", synthetic_source)
        self.assertNotIn("def _clone", desired_source)
        self.assertNotIn("def _thaw", synthetic_source)
        clients_source = (ROOT / "online_e2e" / "clients.py").read_text(encoding="utf-8")
        runner_source = (ROOT / "online_e2e" / "runner.py").read_text(encoding="utf-8")
        self.assertIn("clone_json(value, allow_tuple=False)", clients_source)
        self.assertIn("clone_json(value, allow_tuple=False)", runner_source)
        self.assertNotIn("def _receipt_value", runner_source)

    def test_fixture_ownership_is_frozen_and_encapsulates_exact_set_operations(self):
        from dataclasses import FrozenInstanceError
        from online_e2e.backends import FixtureOwnership

        target = ("AUTO-RUN-01", "AUTO-RUN-02")
        empty = FixtureOwnership("run", target, ())
        self.assertFalse(empty.collision)
        self.assertEqual(empty.created_from(target), target)
        self.assertEqual(empty.cleanup_case_ids(target), target)
        self.assertTrue(empty.is_restored(()))
        self.assertEqual(
            empty.as_dict(),
            {
                "run_id": "run",
                "target_case_ids": ["AUTO-RUN-01", "AUTO-RUN-02"],
                "baseline_case_ids": [],
            },
        )
        collision = FixtureOwnership("run", target, ("AUTO-RUN-01",))
        self.assertTrue(collision.collision)
        with self.assertRaises(FrozenInstanceError):
            empty.run_id = "changed"

        runner_source = (ROOT / "online_e2e" / "runner.py").read_text(encoding="utf-8")
        self.assertNotIn("baseline_case_ids", runner_source)
        self.assertNotIn("fixture_case_ids", runner_source)
        self.assertNotIn('run_state["', runner_source)

    def test_desired_config_rejects_unknown_missing_or_duplicate_skill_before_mutation(self):
        from online_e2e.desired_config import DesiredConfigError, build_desired_config

        variants = []
        unknown = self._online_config()
        unknown["skillInfoList"].append(
            {"skillNid": "unknown", "name": "unknown", "enabled": True}
        )
        variants.append(unknown)
        missing = self._online_config()
        missing["skillInfoList"].pop()
        variants.append(missing)
        duplicate = self._online_config()
        duplicate["skillInfoList"].append(deepcopy(duplicate["skillInfoList"][0]))
        variants.append(duplicate)

        for current in variants:
            with self.subTest(skills=len(current["skillInfoList"])):
                with self.assertRaises(DesiredConfigError):
                    build_desired_config(current, self.target, "new")

    def test_fixed_student_suite_has_six_named_scenarios(self):
        from online_e2e.fixtures import student_scenarios

        scenarios = student_scenarios()
        self.assertEqual(len(scenarios), 6)
        self.assertEqual(
            [item.scenario_id for item in scenarios],
            [
                "exact-statute",
                "detailed-explanation",
                "follow-up-question",
                "ambiguous-candidates",
                "unknown-case-no-fabrication",
                "student-write-denied",
            ],
        )

        by_id = {item.scenario_id: item for item in scenarios}
        self.assertEqual(by_id["exact-statute"].expected_outcome, "answered")
        self.assertTrue(by_id["exact-statute"].needs_reflection)
        self.assertEqual(by_id["detailed-explanation"].required_evidence[0].field, "facts")
        self.assertEqual(
            by_id["detailed-explanation"].any_true_evidence,
            (("dispute", "analysis"),),
        )
        self.assertGreaterEqual(by_id["ambiguous-candidates"].min_candidates, 2)
        self.assertIn("facts", by_id["unknown-case-no-fabrication"].forbidden_fields)
        self.assertFalse(by_id["student-write-denied"].expected_write_performed)
        self.assertEqual({item.continuation_group for item in scenarios}, {"full-suite"})

        runner_source = (ROOT / "online_e2e" / "runner.py").read_text(encoding="utf-8")
        synthetic_source = (ROOT / "online_e2e" / "synthetic_backend.py").read_text(encoding="utf-8")
        self.assertNotIn('scenario_id == "', runner_source)
        self.assertNotIn("receipts = {", synthetic_source)

    def test_teacher_docx_contains_two_explicitly_fictional_cases_and_run_id(self):
        from online_e2e.fixtures import build_teacher_docx, teacher_case_ids

        payload = build_teacher_docx("run_20260904")
        document = Document(BytesIO(payload))
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        text += "\n" + "\n".join(
            cell.text for table in document.tables for row in table.rows for cell in row.cells
        )

        for case_id in teacher_case_ids("run_20260904"):
            self.assertIn(case_id, text)
        self.assertIn("FICTIONAL TEST CASES", text)
        self.assertIn("NO REAL PII", text)
        with ZipFile(BytesIO(payload)) as archive:
            document_xml = archive.read("word/document.xml").decode("utf-8")
        self.assertIn('w:eastAsia="Arial Unicode MS"', document_xml)
        self.assertIn('w:before="160"', document_xml)

    def test_run_id_rejects_path_traversal_for_all_fixture_paths(self):
        from online_e2e.fixtures import build_teacher_docx, teacher_case_ids

        for value in ("../escape", "two/levels", "..", "", "含中文"):
            with self.subTest(value=value):
                with self.assertRaises(ValueError):
                    teacher_case_ids(value)
                with self.assertRaises(ValueError):
                    build_teacher_docx(value)

    def test_teacher_fixture_renders_to_page_png_with_bundled_document_runtime(self):
        from online_e2e.fixtures import (
            assert_rendered_pages_visible,
            render_teacher_fixture_for_qa,
        )
        from PIL import Image

        python_value = os.environ.get("POLYMAS_QA_PYTHON")
        renderer_value = os.environ.get("POLYMAS_DOCX_RENDERER")
        if not python_value or not renderer_value:
            self.skipTest("DOCX 视觉验收需显式注入 POLYMAS_QA_PYTHON 和 POLYMAS_DOCX_RENDERER；结构测试仍执行")
        python, renderer = Path(python_value), Path(renderer_value)
        self.assertTrue(python.is_file(), "POLYMAS_QA_PYTHON 不存在")
        self.assertTrue(renderer.is_file(), "POLYMAS_DOCX_RENDERER 不存在")
        with tempfile.TemporaryDirectory() as temporary:
            docx_path, pages = render_teacher_fixture_for_qa(
                "run_render_001",
                Path(temporary),
                python_executable=python,
                renderer=renderer,
            )
            self.assertTrue(docx_path.is_file())
            self.assertGreaterEqual(len(pages), 1)
            self.assertTrue(all(page.is_file() and page.stat().st_size > 0 for page in pages))
            assert_rendered_pages_visible(pages)

            blank = Path(temporary) / "blank.png"
            Image.new("RGB", (800, 1000), "white").save(blank)
            with self.assertRaisesRegex(RuntimeError, "blank_rendered_page"):
                assert_rendered_pages_visible((blank,))


if __name__ == "__main__":
    unittest.main()
