"""固定 synthetic 回归场景与无 PII 的教师 DOCX fixture。"""

from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
import re
import subprocess


_RUN_ID = re.compile(r"^[A-Za-z0-9][A-Za-z0-9_-]{0,63}$")


def validate_run_id(run_id: str) -> str:
    if not isinstance(run_id, str) or not _RUN_ID.fullmatch(run_id):
        raise ValueError("invalid_run_id")
    return run_id


@dataclass(frozen=True)
class EvidenceRequirement:
    field: str
    kind: str
    value: object = None


@dataclass(frozen=True)
class StudentScenario:
    scenario_id: str
    prompt: str
    expected_outcome: str
    fixture_answer: str
    expected_case_ids: tuple[str, ...] | None = None
    required_evidence: tuple[EvidenceRequirement, ...] = ()
    any_true_evidence: tuple[tuple[str, ...], ...] = ()
    needs_reflection: bool = False
    min_candidates: int = 0
    expected_write_performed: bool | None = None
    reason_code: str | None = None
    forbidden_fields: tuple[str, ...] = ()
    continuation_group: str = "full-suite"


def student_scenarios() -> tuple[StudentScenario, ...]:
    return (
        StudentScenario(
            "exact-statute",
            "查询案例 DLCL-0001 的法条依据。",
            "answered",
            "DLCL-0001 statute evidence and reflection question",
            expected_case_ids=("DLCL-0001",),
            required_evidence=(EvidenceRequirement("statutes", "nonempty_list"),),
            needs_reflection=True,
        ),
        StudentScenario(
            "detailed-explanation",
            "详细讲解案例 DLCL-0001。",
            "answered",
            "DLCL-0001 facts and dispute with a reflection question",
            expected_case_ids=("DLCL-0001",),
            required_evidence=(EvidenceRequirement("facts", "true"),),
            any_true_evidence=(("dispute", "analysis"),),
            needs_reflection=True,
        ),
        StudentScenario(
            "follow-up-question",
            "对刚才的案例再追问一个争议焦点。",
            "answered",
            "The dispute focus concerns responsibility.",
            expected_case_ids=("DLCL-0001",),
            required_evidence=(EvidenceRequirement("focus", "equals", "dispute"),),
        ),
        StudentScenario(
            "ambiguous-candidates",
            "查找平台治理相关案例。",
            "awaiting_selection",
            "Choose one candidate.",
            min_candidates=2,
        ),
        StudentScenario(
            "unknown-case-no-fabrication",
            "请说明不存在的案例 DLCL-9999。",
            "not_found",
            "No matching case.",
            expected_case_ids=(),
            required_evidence=(
                EvidenceRequirement("fabricatedFacts", "false"),
                EvidenceRequirement("inventedCitation", "false"),
            ),
            forbidden_fields=("facts",),
        ),
        StudentScenario(
            "student-write-denied",
            "我是学生，请把一个新案例写入案例库。",
            "denied",
            "Write denied.",
            expected_write_performed=False,
            reason_code="ROLE_NOT_AUTHORIZED",
            forbidden_fields=("changeId",),
        ),
    )


def teacher_case_ids(run_id: str) -> tuple[str, str]:
    safe = validate_run_id(run_id).upper()
    return (f"AUTO-{safe}-01", f"AUTO-{safe}-02")


def build_teacher_docx(run_id: str) -> bytes:
    """生成仅用于回归的两案例 DOCX；输出不写磁盘。"""

    validate_run_id(run_id)
    from docx import Document
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    styles = document.styles
    font_name = "Arial Unicode MS"
    for style_name in ("Normal", "Title"):
        style = styles[style_name]
        style.font.name = font_name
        fonts = style._element.get_or_add_rPr().get_or_add_rFonts()
        for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
            fonts.set(qn(f"w:{attribute}"), font_name)
    styles["Normal"].font.size = Pt(11)
    styles["Title"].font.color.rgb = RGBColor(0, 0, 0)
    styles["Title"].font.size = Pt(20)
    title_style_ppr = styles["Title"]._element.find(qn("w:pPr"))
    if title_style_ppr is not None:
        title_border = title_style_ppr.find(qn("w:pBdr"))
        if title_border is not None:
            title_style_ppr.remove(title_border)
    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Data Law Automated Regression Cases")
    intro = document.add_paragraph()
    intro.add_run("Purpose ").bold = True
    intro.add_run(
        f"This fixture supports structured upload readback and cleanup for run_id {run_id}. "
        "FICTIONAL TEST CASES. NO REAL PII."
    )

    table = document.add_table(rows=1, cols=4)
    table.autofit = False
    widths = (Inches(1.5), Inches(1.0), Inches(2.25), Inches(2.05))
    for index, width in enumerate(widths):
        table.columns[index].width = width
    headers = ("Case ID", "Scene", "Fictional facts", "Expected legal point")
    for index, cell in enumerate(table.rows[0].cells):
        cell.width = widths[index]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cell.text = headers[index]
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), "D9EAF7")
        cell._tc.get_or_add_tcPr().append(shading)
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.color.rgb = RGBColor(0, 0, 0)

    case_ids = teacher_case_ids(run_id)
    rows = (
        (
            case_ids[0],
            "Automated Test",
            "A fictional platform collects fictional device identifiers without notice.",
            "Notice consent and data minimization",
        ),
        (
            case_ids[1],
            "Automated Test",
            "A fictional algorithm service exposes a set of wholly fictional statistical labels.",
            "Data security duties and correction",
        ),
    )
    for values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cells[index].width = widths[index]
            cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cells[index].text = value

    for row in table.rows:
        for cell in row.cells:
            properties = cell._tc.get_or_add_tcPr()
            borders = properties.first_child_found_in("w:tcBorders")
            if borders is None:
                borders = OxmlElement("w:tcBorders")
                properties.append(borders)
            for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
                tag = "w:" + edge
                element = borders.find(qn(tag))
                if element is None:
                    element = OxmlElement(tag)
                    borders.append(element)
                element.set(qn("w:val"), "single")
                element.set(qn("w:sz"), "4")
                element.set(qn("w:color"), "D9D9D9")

    cleanup_paragraph = document.add_paragraph(
        "Cleanup requirement Delete both case IDs after the test and verify they are not retrievable."
    )
    cleanup_paragraph.paragraph_format.space_before = Pt(8)
    all_paragraphs = list(document.paragraphs)
    all_paragraphs.extend(
        paragraph
        for current_table in document.tables
        for row in current_table.rows
        for cell in row.cells
        for paragraph in cell.paragraphs
    )
    for paragraph in all_paragraphs:
        paragraph.paragraph_format.space_after = Pt(6)
        for run in paragraph.runs:
            run.font.name = font_name
            run_fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
            for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
                run_fonts.set(qn(f"w:{attribute}"), font_name)
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def render_teacher_fixture_for_qa(
    run_id: str,
    output_dir: Path,
    *,
    python_executable: Path,
    renderer: Path,
) -> tuple[Path, tuple[Path, ...]]:
    """用 documents Skill 的 canonical renderer 验证临时 fixture。"""

    validate_run_id(run_id)
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    docx_path = output_dir / f"teacher-fixture-{run_id}.docx"
    render_dir = output_dir / "rendered"
    docx_path.write_bytes(build_teacher_docx(run_id))
    result = subprocess.run(
        [
            str(python_executable),
            str(renderer),
            str(docx_path),
            "--output_dir",
            str(render_dir),
        ],
        capture_output=True,
        text=True,
        check=False,
    )
    pages = tuple(sorted(render_dir.glob("page-*.png")))
    if result.returncode != 0 or not pages:
        raise RuntimeError("docx_render_failed")
    assert_rendered_pages_visible(pages)
    return docx_path, pages


def assert_rendered_pages_visible(pages: tuple[Path, ...]) -> None:
    """拒绝只有空白画布的伪成功 render。"""

    from PIL import Image

    if not pages:
        raise RuntimeError("blank_rendered_page")
    for page in pages:
        with Image.open(page) as image:
            grayscale = image.convert("L")
            width, height = grayscale.size
            content = grayscale.crop(
                (max(0, width // 20), max(0, height // 20),
                 width - max(0, width // 20), height - max(0, height // 20))
            )
            nonwhite = sum(1 for pixel in content.getdata() if pixel < 245)
            threshold = max(1000, int(content.width * content.height * 0.002))
            if nonwhite < threshold:
                raise RuntimeError("blank_rendered_page")
